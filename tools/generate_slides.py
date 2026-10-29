"""Generate a five-minute deck directly from the isolated benchmark results."""
import json
from pathlib import Path
import shutil
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from bench.analyze_isolated import validate, stats, BASE
from bench.runner import ROOT


def build_all_slides():
    rows, _ = validate()
    counts = stats(rows)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    def slide(title, bullets, figure=None):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(248,250,252)
        box = s.shapes.add_textbox(Inches(.55), Inches(.35), Inches(12.2), Inches(.85))
        r = box.text_frame.paragraphs[0].add_run(); r.text=title
        r.font.size=Pt(29); r.font.bold=True; r.font.color.rgb=RGBColor(15,23,42)
        width=5.1 if figure else 11.9
        box=s.shapes.add_textbox(Inches(.65), Inches(1.5), Inches(width), Inches(5.25))
        box.text_frame.word_wrap=True
        for index, text in enumerate(bullets):
            p=box.text_frame.paragraphs[0] if index==0 else box.text_frame.add_paragraph()
            p.text=text; p.font.size=Pt(21 if figure else 23); p.space_after=Pt(19)
            p.font.color.rgb=RGBColor(51,65,85)
        if figure:
            s.shapes.add_picture(str(BASE/'figures'/figure), Inches(5.95), Inches(1.85), width=Inches(6.85))
        footer=s.shapes.add_textbox(Inches(.65), Inches(7.0), Inches(12), Inches(.3))
        p=footer.text_frame.paragraphs[0]; p.text=f'CSD203 · Nhóm 4 · AI21B · {len(prs.slides)} / 8';p.font.size=Pt(10)
        return s
    slide('Nén dữ liệu và Byte Pair Encoding', [
        'Ba yếu tố: kích thước dữ liệu, thứ tự byte và corpus huấn luyện.',
        'RLE · Canonical Huffman · LZW · BPE naive / noheap / inc',
        'Nhóm 4 — AI21B — FA26',
        '233 cấu hình; số liệu mới từ các tiến trình độc lập.'])
    slide('Giao thức đo và kiểm chứng', [
        'Mỗi cấu hình: 3 repetition, seed 1/2/3, tiến trình mới, timeout 60 giây.',
        'GC mặc định bật; perf_counter_ns chỉ đo thuật toán.',
        'Bộ nhớ: tiến trình --mem riêng; timeout được ghi riêng, không thay bằng 0.',
        'Full SHA-256, CSV median + min/max, nhật ký từng tiến trình.'])
    slide('RQ1 — Quy mô và phân bố dữ liệu', [
        '5 kích thước: 256–4096 KiB; 150 cấu hình.',
        'C1: run dài; C2: từ Zipf; C3: byte ngẫu nhiên; V1: văn bản Việt.',
        'RLE hưởng lợi từ run dài. Byte ngẫu nhiên có thể làm dữ liệu nén lớn hơn đầu vào.',
        'Dải trên hình là biến thiên giữa seed input.'], 'rq1_scaling.png')
    slide('RQ2 — Cùng histogram, khác cấu trúc chuỗi', [
        '45 cấu hình: gốc, gom run, shuffle 3 seed.',
        'Huffman giữ nguyên kích thước container khi histogram không đổi.',
        'RLE và LZW thay đổi vì khai thác run và chuỗi con.',
        'Gom run chưa bao gồm chi phí lưu hoán vị ngược.'], 'rq2_order.png')
    slide('RQ3 — Corpus huấn luyện ảnh hưởng số token', [
        'Train/test V1 và E1: ma trận 2×2, full corpus.',
        'K = 256 / 1024 / 3840; ratio gồm bảng merge và header.',
        'Chuyển corpus có thể tăng số token trên cùng input.',
        'Corpus không song song: không suy rộng thành bản chất của ngôn ngữ.'], 'rq3_ratio.png')
    slide('Ablation — Giá trị của cập nhật và heap', [
        'naive: đếm lại toàn bộ cặp mỗi merge.',
        'noheap: cập nhật gia tăng, quét để chọn cặp.',
        'inc: cập nhật gia tăng + heap lazy.',
        'Mỗi biến thể có 3 lần thử; timeout được báo riêng, không tính speedup từ số chưa đo được.'], 'rq3_training.png')
    summary=[f"RQ{rq}: {c['ok']} ok, {c['timeout']} timeout, {c['crash']+c['oom']+c['wrong']} lỗi khác." for rq,c in counts.items()]
    slide('Kết quả đáng chú ý và giới hạn', summary+[
        'Cùng entropy bậc 0 không có nghĩa cùng tỷ lệ nén với codec khai thác thứ tự.',
        'Memory tracing có thể timeout dù lượt đo thời gian thành công.',
        'Ba repetition và một máy đo chưa đủ để suy rộng mọi môi trường.'])
    slide('Kết luận và demo tái lập', [
        'Chọn thuật toán theo cấu trúc input và corpus; tính cả chi phí học BPE.',
        'python -m pytest -q',
        'python -m bench.runner --verify results/isolated/results_rq3.csv --row 8',
        'Đối chiếu status/hash chính xác; thời gian trong 0.5–2× median.',
        'AI hỗ trợ sửa mã và báo cáo; sinh viên tự review và bảo vệ phần việc.'])
    target=ROOT/'paper/CSD203_DA4_Presentation.pptx'
    prs.save(target)
    shutil.copy2(target, ROOT/'final_submission/CSD203_DA4_Presentation.pptx')
    print(f'Built {len(prs.slides)} slides from isolated results')


if __name__=='__main__': build_all_slides()
