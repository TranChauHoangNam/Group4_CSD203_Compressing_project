"""Check DOCX conversion, source provenance and removal of superseded reports."""
import json
import re
from docx import Document
from docx.oxml.ns import qn

from final_submission.generate_report import OUT, TARGET, report, sources
from final_submission.build_submission import FIGS


def main():
    text=TARGET.read_text(encoding='utf-8')
    assert text==report()
    assert json.loads((OUT/'report_sources.json').read_text())==sources()
    document=Document(OUT/'CSD203_DA4_Report_Revised.docx')
    expected=[]
    current=[]
    for line in text.splitlines()+['']:
        if line.startswith('|'):
            row=[v.strip() for v in line.strip('|').split('|')]
            if not all(re.fullmatch(r':?-{3,}:?',v) for v in row):
                current.append(row)
        elif current:
            expected.append(current);current=[]
    actual=[[[c.text for c in row.cells] for row in table.rows] for table in document.tables]
    assert actual==expected and len(actual)==6, 'DOCX tables differ from report'
    paragraphs='\n'.join(p.text for p in document.paragraphs)
    for member in json.loads((OUT/'team.json').read_text(encoding='utf-8'))['members']:
        assert member in paragraphs
    assert '\ufffd' not in paragraphs and '104.72%' not in paragraphs
    assert len(document.inline_shapes)==7
    assert any(s._sectPr.find(qn('w:cols')).get(qn('w:num'))=='2' for s in document.sections)
    for name,source in FIGS.items():
        assert (OUT/'figures'/name).read_bytes()==source.read_bytes(), f'Stale figure: {name}'
    assert not list(OUT.glob('CSD203_DA4_IEEE_Final.*')), 'Superseded report still present'
    print('PASS: six DOCX tables, six authors, seven current figures, two-column sections and no superseded reports')


if __name__=='__main__': main()
