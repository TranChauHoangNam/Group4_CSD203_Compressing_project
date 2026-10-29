"""Build DOCX and local figure bundle from the final Markdown report."""
from pathlib import Path
import re
import shutil
from html import escape

from docx import Document
from docx.image.image import Image as DocxImage
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
MD = OUT / "CSD203_DA4_Report_Revised.md"
FIGS = {
    "rq1_ratio.png": ROOT / "rq1/out/rq1_ratio.png",
    "rq1_throughput.png": ROOT / "rq1/out/rq1_throughput.png",
    "rq1_memory.png": ROOT / "rq1/out/rq1_memory.png",
    "rq2_ratio.png": ROOT / "rq2/out/rq2_ratio.png",
    "rq3_bytes_per_token.png": ROOT / "rq3/out/rq3_bytes_per_token.png",
    "rq3_train_ms.png": ROOT / "rq3/out/rq3_train_ms.png",
    "rq3_training_memory.png": ROOT / "rq3/out/rq3_training_memory.png",
}


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, value, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(value.replace("`", "").replace("**", ""))
    run.font.name = "Times New Roman"
    run.font.size = Pt(8.5)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("CSD203 · Đề tài 4  |  ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def format_inline(paragraph, text, size=10):
    # Keep the report readable without depending on external Markdown packages.
    chunks = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for chunk in chunks:
        if not chunk:
            continue
        bold = chunk.startswith("**") and chunk.endswith("**")
        code = chunk.startswith("`") and chunk.endswith("`")
        value = chunk[2:-2] if bold else chunk[1:-1] if code else chunk
        run = paragraph.add_run(value)
        run.bold = bold
        run.font.name = "Consolas" if code else "Times New Roman"
        run.font.size = Pt(size)


def html_inline(text):
    value = escape(text)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", value)
    value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
    value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"<em>\1</em>", value)
    return value


def build():
    text = MD.read_text(encoding="utf-8")
    if "Discussion skeleton" in text or "GIAI THICH - can kiem chung" in text:
        raise SystemExit("Unresolved placeholder in final report")
    (OUT / "figures").mkdir(exist_ok=True)
    for name, src in FIGS.items():
        if not src.is_file() or src.stat().st_size == 0:
            raise FileNotFoundError(src)
        shutil.copy2(src, OUT / "figures" / name)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    def columns(count):
        current = doc.add_section(WD_SECTION.CONTINUOUS)
        element = current._sectPr.find(qn('w:cols'))
        if element is None:
            element = OxmlElement('w:cols')
            current._sectPr.append(element)
        element.set(qn('w:num'), str(count))
        element.set(qn('w:space'), '360')

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for style_name, size, color in (("Heading 1", 10, (0, 0, 0)),
                                    ("Heading 2", 10, (0, 0, 0)),
                                    ("Heading 3", 10, (0, 0, 0))):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    add_page_field(section.footer.paragraphs[0])

    lines = text.splitlines()
    i = 0
    body_started = False
    wide_block = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if wide_block and not line.startswith(('*Hình', '|')):
            columns(2)
            wide_block = False
        if line.startswith('## ') and not body_started:
            columns(2)
            body_started = True
        if line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if match:
                columns(1)
                wide_block = True
                path = OUT / match.group(2)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(0)
                raster = DocxImage.from_file(str(path))
                width = min(7.0, 5.1 * raster.px_width / raster.px_height)
                p.add_run().add_picture(str(path), width=Inches(width))
                p.paragraph_format.keep_with_next = True
            i += 1
            continue
        if line.startswith("|"):
            if not wide_block:
                columns(1)
            wide_block = False
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in row):
                    rows.append(row)
                i += 1
            if rows:
                cols = max(map(len, rows))
                table = doc.add_table(rows=len(rows), cols=cols)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = "Table Grid"
                table.autofit = False
                for column in table.columns:
                    column.width = Inches(7.0 / cols)
                header_properties = table.rows[0]._tr.get_or_add_trPr()
                header_properties.append(OxmlElement('w:tblHeader'))
                for r_idx, row in enumerate(rows):
                    for c_idx in range(cols):
                        value = row[c_idx] if c_idx < len(row) else ""
                        set_cell_text(table.cell(r_idx, c_idx), value,
                                      bold=(r_idx == 0),
                                      color=(0, 0, 0) if r_idx == 0 else None)
                        if r_idx == 0:
                            shade(table.cell(r_idx, c_idx), "EEEEEE")
                doc.add_paragraph().paragraph_format.space_after = Pt(1)
            columns(2)
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:], level=3)
        elif line.startswith("# "):
            title = line[2:]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(title)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith(("* ", "- ")):
            p = doc.add_paragraph(style="List Bullet")
            format_inline(p, line[2:])
        elif re.match(r"\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            format_inline(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith(("*Hình", "*Bảng")):
            if line.startswith('*Bảng'):
                columns(1)
                wide_block = True
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(5)
            if line.startswith('*Bảng'):
                p.paragraph_format.keep_with_next = True
            r = p.add_run(line.strip("*"))
            r.italic = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(8.5)
        elif line.startswith("**"):
            p = doc.add_paragraph()
            format_inline(p, line)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = doc.add_paragraph()
            format_inline(p, line)
        i += 1

    props = doc.core_properties
    props.title = "CSD203 Đề tài 4 — Báo cáo IEEE"
    props.subject = "Nén không mất mát và tokenizer BPE"
    props.author = "Nhóm AI21B — FA26"
    target = OUT / "CSD203_DA4_Report_Revised.docx"
    doc.save(target)
    print(f"WROTE {target} ({target.stat().st_size} bytes)")

    # Produce a print stylesheet twin of the same final Markdown for Edge PDF export.
    html_parts = ["""<!doctype html><html lang=\"vi\"><head><meta charset=\"utf-8\"><style>
@page { size: Letter; margin: 0.5in 0.55in; }
body { font: 8.8pt/1.12 'Times New Roman',serif; color:#000; column-count:2; column-gap:0.2in; }
h1 { font-size:14pt; text-align:center; color:#000; margin:4pt 0 4pt; column-span:all; }
h2 { font-size:9.5pt; color:#000; margin:6pt 0 3pt; break-after:avoid; }
h3 { font-size:9pt; color:#000; margin:5pt 0 2pt; break-after:avoid; }
h4 { font-size:8.8pt; margin:4pt 0 2pt; break-after:avoid; }
p { margin:2.5pt 0; text-align:justify; orphans:2; widows:2; }
table { border-collapse:collapse; width:100%; margin:4pt 0; font-size:7pt; break-inside:avoid; }
th,td { border:0.5pt solid #aab4bd; padding:1.5pt 2.5pt; vertical-align:middle; }
th { background:#eee; color:black; }
img { display:block; width:100%; max-width:100%; max-height:1.65in; object-fit:contain; margin:3pt auto 1pt; break-inside:avoid; }
.caption { text-align:center; font-size:7.8pt; font-style:italic; margin:1pt 0 4pt; break-inside:avoid; }
.author { text-align:center; margin:1.5pt 0; font-size:8.8pt; column-span:all; }
ul,ol { margin:2.5pt 0; padding-left:16pt; } li { margin:1.5pt 0; }
code { font-family:Consolas,monospace; font-size:8pt; }
@media print { a { color:inherit; text-decoration:none; } }
</style></head><body>"""]
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in row):
                    rows.append(row)
                i += 1
            html_parts.append("<table>")
            for r_idx, row in enumerate(rows):
                html_parts.append("<tr>" + "".join(
                    f"<{ 'th' if r_idx == 0 else 'td' }>{escape(c)}</{ 'th' if r_idx == 0 else 'td' }>"
                    for c in row) + "</tr>")
            html_parts.append("</table>")
            continue
        image_match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if image_match:
            alt, path = image_match.groups()
            html_parts.append(f'<img src="{escape(path)}" alt="{escape(alt)}">')
        elif line.startswith(("*Hình", "*Bảng")):
            html_parts.append(f'<p class="caption">{escape(line.strip("*"))}</p>')
        elif line.startswith("#### "):
            html_parts.append(f"<h4>{escape(line[5:])}</h4>")
        elif line.startswith("### "):
            html_parts.append(f"<h3>{escape(line[4:])}</h3>")
        elif line.startswith("## "):
            html_parts.append(f"<h2>{escape(line[3:])}</h2>")
        elif line.startswith("# "):
            html_parts.append(f"<h1>{escape(line[2:])}</h1>")
        elif line.startswith('**') and line.endswith('**'):
            html_parts.append(f'<p class="author">{html_inline(line)}</p>')
        elif re.match(r"[-*] ", line):
            html_parts.append(f"<ul><li>{html_inline(line[2:])}</li></ul>")
        elif re.match(r"\d+\. ", line):
            html_parts.append(f"<ol><li>{html_inline(re.sub(r'^\\d+\\. ', '', line))}</li></ol>")
        else:
            html_parts.append(f"<p>{html_inline(line)}</p>")
        i += 1
    html_parts.append("</body></html>")
    html_path = OUT / "CSD203_DA4_Report_Revised.html"
    html_path.write_text("\n".join(html_parts), encoding="utf-8")
    print(f"WROTE {html_path} ({html_path.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
